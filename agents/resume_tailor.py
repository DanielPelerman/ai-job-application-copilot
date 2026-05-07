import io
import re
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    Document = None
    DOCX_AVAILABLE = False

try:
    from docx2pdf import convert
    DOCX2PDF_AVAILABLE = True
except ImportError:
    convert = None
    DOCX2PDF_AVAILABLE = False


class ResumeTailor:
    """Rule-based resume tailoring generator using resume text, job details, and profile data."""

    KEYWORD_GROUPS = {
        "technical": [
            "python", "sql", "excel", "tableau", "r", "pandas", "numpy", "scikit-learn",
            "streamlit", "visualization", "dashboard", "reporting", "modeling",
        ],
        "finance": [
            "financial modeling", "forecasting", "budgeting", "valuation", "investment", "portfolio",
            "risk", "cash flow", "variance analysis", "p&l", "balance sheet", "econometric",
        ],
        "leadership": [
            "lead", "managed", "mentored", "coordinated", "team", "stakeholder", "presented",
        ],
        "tools": [
            "github", "jupyter", "quarto", "powershell", "aws", "sql server", "excel", "tableau",
        ],
    }

    def suggest(self, resume_text: str, job_description: str) -> dict:
        suggestions = []
        fields_review = []
        lower_resume = resume_text.lower() if resume_text else ""
        lower_job = job_description.lower() if job_description else ""

        if not resume_text:
            suggestions.append("Upload a resume to get focused tailoring suggestions.")
            fields_review.append("Resume content")

        for group in self.KEYWORD_GROUPS.values():
            for skill in group:
                if skill in lower_job and skill not in lower_resume:
                    suggestions.append(f"Consider adding or emphasizing experience with '{skill}' in your resume.")

        if "project" in lower_job and "project" not in lower_resume:
            suggestions.append("Include specific project examples that illustrate your experience.")

        if any(token in lower_job for token in ["lead", "manage", "mentor"]) and all(token not in lower_resume for token in ["lead", "managed", "mentored"]):
            suggestions.append("Highlight leadership, team coordination, or stakeholder collaboration experience.")

        if not lower_job:
            fields_review.append("Job description details")

        if not suggestions:
            suggestions.append(
                "Your resume appears to cover many job keywords; review for clarity, metrics, and business impact."
            )

        if "location" not in lower_job and "remote" not in lower_job:
            fields_review.append("Location or work arrangement")

        if "company" not in lower_job:
            fields_review.append("Company-specific context")

        return {
            "suggestions": suggestions,
            "fields_review": list(dict.fromkeys(fields_review)),
        }

    def _normalize_keywords(self, text: str) -> set:
        words = re.findall(r"[a-zA-Z0-9\+\-]+", text.lower())
        return set(words)

    def _extract_profile_keywords(self, profile: dict) -> set:
        text_parts = []
        text_parts.append(profile.get("technical_skills", ""))
        text_parts.append(profile.get("finance_analytics_skills", ""))
        text_parts.append(profile.get("tools_platforms", ""))
        text_parts.append(profile.get("target_roles", ""))
        text_parts.append(profile.get("work_experience", ""))
        text_parts.append(profile.get("projects", ""))
        text_parts.append(profile.get("resume_bullets", ""))
        return self._normalize_keywords("\n".join(text_parts))

    def _extract_job_keywords(self, job_description: str) -> set:
        lower_description = job_description.lower()
        keywords = set()
        for group in self.KEYWORD_GROUPS.values():
            for keyword in group:
                if keyword in lower_description:
                    keywords.add(keyword)

        if keywords:
            return keywords

        words = re.findall(r"[a-zA-Z0-9\+\-]+", lower_description)
        stop_words = {
            "the", "and", "for", "with", "that", "this", "from", "your",
            "will", "have", "business", "role", "experience", "company",
            "position", "team", "work", "skills", "including", "using",
            "data", "analysis",
        }
        filtered = [w for w in words if w not in stop_words and len(w) > 2]
        frequencies = {}
        for word in filtered:
            frequencies[word] = frequencies.get(word, 0) + 1
        top_terms = sorted(frequencies, key=frequencies.get, reverse=True)[:15]
        return set(top_terms)

    def _extract_project_lines(self, project_text: str, job_keywords: set) -> list:
        lines = [line.strip() for line in project_text.splitlines() if line.strip()]
        highlighted = []
        for line in lines:
            lower_line = line.lower()
            if any(keyword in lower_line for keyword in job_keywords):
                highlighted.append(line)
        return highlighted[:5]

    def _build_tailored_intro(self, matched_keywords: list, highlighted_experience: list, profile: dict, role_title: str, company: str) -> str:
        intro_lines = [f"Tailored resume for {role_title} at {company}", ""]
        if matched_keywords:
            intro_lines.append("Relevant keywords matched:")
            for keyword in matched_keywords:
                intro_lines.append(f"• {keyword}")
            intro_lines.append("")

        if highlighted_experience:
            intro_lines.append("Key experience and projects to emphasize:")
            for item in highlighted_experience:
                intro_lines.append(f"• {item}")
            intro_lines.append("")

        if profile.get("work_experience"):
            intro_lines.append("Preserved core experience from your original resume while emphasizing the most relevant finance and analytics results.")
            intro_lines.append("")

        return "\n".join(intro_lines)

    def _score_match(self, matched_keywords: set, job_keywords: set, profile_keywords: set) -> int:
        if not job_keywords:
            return 0
        base_score = int(min(100, (len(matched_keywords) / max(1, len(job_keywords))) * 100))
        overlap = matched_keywords.intersection(profile_keywords)
        bonus = min(20, len(overlap) * 2)
        return min(100, base_score + bonus)

    def _sanitize_filename(self, value: str) -> str:
        sanitized = re.sub(r"[^a-zA-Z0-9_\-]+", "_", value.strip().lower())
        sanitized = re.sub(r"_+", "_", sanitized)
        return sanitized.strip("_")[:120]

    def _export_docx(self, document: Document, output_path: Path) -> Path:
        if not DOCX_AVAILABLE or document is None:
            return None
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
        return output_path

    def _convert_docx_to_pdf(self, docx_path: Path, pdf_path: Path) -> Path:
        if not DOCX2PDF_AVAILABLE or docx_path is None or not docx_path.exists():
            return None
        try:
            pdf_path.parent.mkdir(parents=True, exist_ok=True)
            convert(str(docx_path), str(pdf_path))
            return pdf_path if pdf_path.exists() else None
        except Exception:
            return None

    def _bold_matched_keywords(self, document: Document, keywords: list) -> None:
        if not keywords or document is None:
            return
        safe_keywords = sorted(set(keywords), key=len, reverse=True)
        pattern = re.compile(r"\b(" + "|".join(re.escape(k) for k in safe_keywords) + r")\b", flags=re.IGNORECASE)

        def bold_runs(runs):
            for run in runs:
                if run.text and pattern.search(run.text):
                    run.bold = True

        for paragraph in document.paragraphs:
            bold_runs(paragraph.runs)

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        bold_runs(paragraph.runs)

    def generate_tailored_resume(
        self,
        resume_docx_bytes: bytes,
        resume_text: str,
        job_description: str,
        profile: dict,
        company: str = "Company",
        role_title: str = "Role",
        original_filename: str = "resume.docx",
    ) -> dict:
        job_keywords = self._extract_job_keywords(job_description)
        profile_keywords = self._extract_profile_keywords(profile)
        resume_keywords = self._normalize_keywords(resume_text)
        matched_keywords = sorted({kw for kw in job_keywords if kw in profile_keywords or kw in resume_keywords})
        missing_qualifications = sorted(
            [kw for kw in job_keywords if kw not in profile_keywords and kw not in resume_keywords]
        )

        emphasized_projects = self._extract_project_lines(profile.get("projects", ""), job_keywords)
        if not emphasized_projects and profile.get("projects"):
            project_lines = [line.strip() for line in profile.get("projects", "").splitlines() if line.strip()]
            emphasized_projects = project_lines[:3]

        recommended_edits = []
        if matched_keywords:
            recommended_edits.append(
                "Emphasize the keywords already present in your resume by preserving and highlighting those exact phrases."
            )
        if emphasized_projects:
            recommended_edits.append("Keep the most relevant project bullets visible and easy to scan.")
        if missing_qualifications:
            recommended_edits.append(
                "If you have related experience, make sure it is visible in the same sections rather than inventing new content."
            )
        if not recommended_edits:
            recommended_edits.append("Your resume already contains strong alignment; preserve the original format and content.")

        match_score = self._score_match(set(matched_keywords), job_keywords, profile_keywords)
        highlighted_experience = self._extract_project_lines(profile.get("projects", ""), job_keywords)
        if not highlighted_experience:
            highlighted_experience = [line for line in profile.get("projects", "").splitlines() if line.strip()][:3]

        document = Document(io.BytesIO(resume_docx_bytes)) if DOCX_AVAILABLE else None
        if document is not None:
            self._bold_matched_keywords(document, matched_keywords)

        date_suffix = datetime.utcnow().strftime("%Y-%m-%d")
        company_name = company or ""
        role_name = role_title or ""
        if company_name.strip() and role_name.strip():
            base = f"{company_name}_{role_name}_{date_suffix}_tailored_resume"
        else:
            base = f"tailored_resume_{date_suffix}"
        sanitized_base = self._sanitize_filename(base)
        output_dir = Path("exports/tailored_resumes")
        docx_path = output_dir / f"{sanitized_base}.docx"
        pdf_path = output_dir / f"{sanitized_base}.pdf"

        actual_docx_path = self._export_docx(document, docx_path) if document is not None else None
        actual_pdf_path = self._convert_docx_to_pdf(actual_docx_path, pdf_path) if actual_docx_path is not None else None

        docx_bytes = None
        pdf_bytes = None
        if actual_docx_path and actual_docx_path.exists():
            docx_bytes = actual_docx_path.read_bytes()
        if actual_pdf_path and actual_pdf_path.exists():
            pdf_bytes = actual_pdf_path.read_bytes()

        result = {
            "matched_keywords": matched_keywords,
            "emphasized_projects": emphasized_projects,
            "recommended_edits": recommended_edits,
            "missing_qualifications": missing_qualifications,
            "match_score": match_score,
            "docx_filename": actual_docx_path.name if actual_docx_path else None,
            "pdf_filename": actual_pdf_path.name if actual_pdf_path else None,
            "docx_path": str(actual_docx_path) if actual_docx_path else None,
            "pdf_path": str(actual_pdf_path) if actual_pdf_path else None,
            "docx_bytes": docx_bytes,
            "pdf_bytes": pdf_bytes,
            "directory": str(output_dir),
        }
        return result
